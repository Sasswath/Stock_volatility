#!/usr/bin/env python3
"""Generate a repository-grounded PowerPoint deck for Stock_volatility."""

from __future__ import annotations

import csv
from pathlib import Path
from typing import Dict, List, Tuple

from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "output"
FIG_DIR = OUT_DIR / "figures"
TAB_DIR = OUT_DIR / "tables"
PPT_PATH = OUT_DIR / "Stock_volatility_analysis.pptx"


def read_dict_csv(path: Path) -> List[Dict[str, str]]:
    with path.open("r", newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def read_rep_selection(path: Path) -> Dict[str, str]:
    rows = read_dict_csv(path)
    return {r["field"]: r["value"] for r in rows}


def find_max_correlation(path: Path) -> Tuple[str, str, float]:
    with path.open("r", newline="", encoding="utf-8") as f:
        rows = list(csv.reader(f))
    headers = rows[0][1:]
    best = ("", "", -1.0)
    for i, row in enumerate(rows[1:]):
        r_ticker = row[0].strip('"')
        for j, v in enumerate(row[1:]):
            c_ticker = headers[j].strip('"')
            if r_ticker == c_ticker:
                continue
            val = float(v)
            if val > best[2]:
                best = (r_ticker, c_ticker, val)
    return best


def extract_preprocessing_snippet(path: Path) -> str:
    lines = path.read_text(encoding="utf-8").splitlines()
    start = None
    end = None
    for i, ln in enumerate(lines):
        if "df$log_ret <-" in ln:
            start = max(0, i - 2)
        if start is not None and "write_csv(ext_ret" in ln:
            end = min(len(lines), i + 6)
            break
    if start is None:
        return "# Preprocessing snippet unavailable"
    block = lines[start:end] if end else lines[start : start + 14]
    return "\n".join(block)


def add_title_slide(prs: Presentation, title: str, subtitle: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle


def add_bullets_slide(prs: Presentation, title: str, bullets: List[str], font_size: int = 20) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    tf = slide.shapes.placeholders[1].text_frame
    tf.clear()
    for i, line in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.level = 0
        p.font.size = Pt(font_size)


def add_image_with_bullets(
    prs: Presentation,
    title: str,
    image_path: Path,
    bullets: List[str],
    image_left: float = 0.3,
    image_top: float = 1.3,
    image_width: float = 6.2,
) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = title
    if image_path.exists():
        slide.shapes.add_picture(str(image_path), Inches(image_left), Inches(image_top), width=Inches(image_width))
    textbox = slide.shapes.add_textbox(Inches(6.7), Inches(1.3), Inches(6.3), Inches(5.7))
    tf = textbox.text_frame
    tf.clear()
    for i, line in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.size = Pt(18)
        p.level = 0


def add_data_prep_slide(prs: Presentation, snippet: str, bullets: List[str]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "3) Data preparation"

    code_box = slide.shapes.add_textbox(Inches(0.3), Inches(1.2), Inches(7.8), Inches(5.7))
    code_tf = code_box.text_frame
    code_tf.clear()
    p = code_tf.paragraphs[0]
    p.text = snippet
    p.font.name = "Courier New"
    p.font.size = Pt(11)

    txt_box = slide.shapes.add_textbox(Inches(8.2), Inches(1.2), Inches(4.8), Inches(5.7))
    tf = txt_box.text_frame
    tf.clear()
    for i, line in enumerate(bullets):
        pp = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        pp.text = line
        pp.level = 0
        pp.font.size = Pt(16)


def pct(v: float, digits: int = 2) -> str:
    return f"{v * 100:.{digits}f}%"


def main() -> None:
    # Load repository outputs
    rep_sel = read_rep_selection(TAB_DIR / "representative_stock_selection.csv")
    clean_rows = read_dict_csv(TAB_DIR / "cleaning_summary.csv")
    desc_rows = read_dict_csv(TAB_DIR / "5stock_descriptive_stats.csv")
    model_rows = read_dict_csv(TAB_DIR / "model_comparison_summary.csv")
    gof_rows = read_dict_csv(TAB_DIR / "goodness_of_fit_summary.csv")
    mc_rows = read_dict_csv(TAB_DIR / "monte_carlo_summary.csv")
    var_rows = read_dict_csv(TAB_DIR / "var_table.csv")
    clt_rows = read_dict_csv(TAB_DIR / "clt_summary.csv")
    t_params = read_dict_csv(TAB_DIR / "t_fit_params.csv")[0]

    rep_ticker = rep_sel.get("Chosen representative ticker", "TCS")
    rep_desc = next(r for r in desc_rows if r["Ticker"] == rep_ticker)
    max_corr = find_max_correlation(TAB_DIR / "correlation_matrix.csv")
    snippet = extract_preprocessing_snippet(ROOT / "R" / "03_data_cleaning.R")

    total_raw = sum(int(float(r["n_raw"])) for r in clean_rows)
    total_clean = sum(int(float(r["n_clean"])) for r in clean_rows)
    total_extreme = sum(int(float(r["n_extreme_ret"])) for r in clean_rows)

    model_normal = next(r for r in model_rows if r["Model"] == "Normal")
    model_t = next(r for r in model_rows if "Student" in r["Model"])
    ks_normal = next(r for r in gof_rows if "KS (Normal)" in r["Test"])
    ks_t = next(r for r in gof_rows if "KS (Student-t)" in r["Test"])

    mc_normal = next(r for r in mc_rows if r["Model"] == "Normal")
    mc_t = next(r for r in mc_rows if r["Model"] == "Student-t")
    mc_emp = next(r for r in mc_rows if r["Model"] == "Empirical resampling")

    clt_n100 = next(r for r in clt_rows if r["n"] == "100")

    rep_var_95 = [r for r in var_rows if r["Asset"] == "Representative stock" and r["Confidence"] == "0.95"]

    prs = Presentation()

    # 1
    add_title_slide(
        prs,
        "Stock Volatility Analysis (Sasswath/Stock_volatility)",
        "Dataset: 5 NSE stocks (2020-01-01 to 2025-12-30), 1,485 daily log returns per stock",
    )

    # 2
    add_bullets_slide(
        prs,
        "2) Aim",
        [
            "Model daily stock-return behavior and risk using repository outputs.",
            "Compare Normal vs Student-t fits for representative stock (TCS).",
            "Quantify tails, volatility, simulation outcomes, and VaR implications.",
        ],
    )

    # 3
    add_data_prep_slide(
        prs,
        snippet,
        [
            f"Raw rows: {total_raw}; cleaned rows: {total_clean}; removed: {total_raw - total_clean}.",
            "Cleaning checks: datatypes, duplicates, NA, invalid prices, date ordering.",
            "Outliers/extreme returns are flagged (|z_ret| > 5) and retained for analysis.",
            f"Total flagged extreme returns across 5 stocks: {total_extreme}.",
            "Source: R/03_data_cleaning.R + output/tables/cleaning_summary.csv",
        ],
    )

    # 4
    add_image_with_bullets(
        prs,
        "4) Data visualization",
        FIG_DIR / "historical_prices.png",
        [
            "Historical adjusted prices and cumulative-return patterns differ by stock.",
            "Distribution/volatility visuals are built from output/figures artifacts.",
            "Representative stock selected objectively: TCS.",
        ],
    )

    # 5
    add_image_with_bullets(
        prs,
        "5) Outlier analysis",
        FIG_DIR / "extreme_return_plot.png",
        [
            "Method: return z-scores on daily log returns; extreme if |z_ret| > 5.",
            "Extreme-return counts: HDFCBANK 4, INFOSYS 7, LT 8, RELIANCE 4, TCS 2.",
            "These events are analyzed as financially meaningful tail behavior.",
        ],
    )

    # 6
    add_bullets_slide(
        prs,
        "6) Removing outliers",
        [
            "No automatic outlier deletion was applied in this repository pipeline.",
            "Evidence: n_raw == n_clean for all 5 stocks in cleaning_summary.csv.",
            "Reason: preserving tail events is essential for volatility/tail-risk/VaR analysis.",
            "Retained data supports robust comparison of Normal vs heavy-tailed Student-t.",
        ],
    )

    # 7
    add_image_with_bullets(
        prs,
        "7) Graphical analysis",
        FIG_DIR / "qq_plots_combined.png",
        [
            "Diagnostic plots: Q-Q (Normal vs Student-t), density overlays, CDF comparison.",
            "Representative stock: TCS; sample skewness ≈ -0.282, kurtosis ≈ 7.905.",
            "Visual tail departures motivate heavy-tail modeling.",
        ],
    )

    # 8
    add_image_with_bullets(
        prs,
        "8) Correlation and variance",
        FIG_DIR / "correlation_heatmap.png",
        [
            f"Highest pairwise correlation: {max_corr[0]}–{max_corr[1]} = {max_corr[2]:.3f}.",
            f"TCS annualized volatility: {rep_sel.get('Annualized volatility', 'N/A')}.",
            "Volatility differs by stock (see comparative_volatility.png).",
        ],
    )

    # 9
    add_bullets_slide(
        prs,
        "9) Regression analysis",
        [
            "Unavailable in current repository outputs/scripts.",
            "No regression model (e.g., lm/glm) artifacts are produced for this workflow.",
            "The analysis focuses on distribution fitting, diagnostics, simulation, and VaR.",
            "Result reporting remains transparent: no fabricated regression assumptions or metrics.",
        ],
    )

    # 10
    add_image_with_bullets(
        prs,
        "10) Prediction",
        FIG_DIR / "monte_carlo_simulation.png",
        [
            "Prediction artifacts are simulation-based (not price-forecast models).",
            f"Normal MC: SD={pct(float(mc_normal['SD']))}, P1={pct(float(mc_normal['P1']))}, loss prob={float(mc_normal['Prob_Loss']):.4f}.",
            f"Student-t MC: SD={pct(float(mc_t['SD']))}, P1={pct(float(mc_t['P1']))}, loss prob={float(mc_t['Prob_Loss']):.4f}.",
            f"Empirical resampling: P1={pct(float(mc_emp['P1']))}, loss prob={float(mc_emp['Prob_Loss']):.4f}.",
        ],
    )

    # 11
    add_bullets_slide(
        prs,
        "11) Comparing the model",
        [
            f"Normal: AIC {float(model_normal['AIC']):.2f}, BIC {float(model_normal['BIC']):.2f}, KS p={ks_normal['p_value']}.",
            f"Student-t: AIC {float(model_t['AIC']):.2f}, BIC {float(model_t['BIC']):.2f}, KS p={ks_t['p_value']}.",
            f"Student-t df estimate: {float(t_params['df']):.3f} (heavier tails than Normal).",
            "Interpretation: Student-t fits observed tail behavior better on this dataset.",
            "Why compare these two? Normal is the light-tailed benchmark; Student-t captures excess kurtosis common in returns.",
        ],
        font_size=18,
    )

    # 12
    add_bullets_slide(
        prs,
        "12) Gain curve plot",
        [
            "Not applicable for this repository's return-distribution workflow.",
            "Gain/lift curves require ranked classification with a binary target.",
            "Here, outputs model continuous financial returns, volatility, and tail quantiles (VaR), not class probabilities.",
        ],
    )

    # 13
    add_bullets_slide(
        prs,
        "13) Conclusion",
        [
            f"Dataset: five stocks, {rep_desc['Rows']} daily returns each; representative ticker: {rep_ticker}.",
            f"CLT check at n=100: empirical/theoretical SE ratio = {float(clt_n100['Ratio']):.3f}.",
            "Heavy tails matter: Student-t improves fit versus Normal on TCS tail behavior.",
            "Why no z-test: population variance is unknown; returns are heavy-tailed/dependent, so known-variance mean-z testing is not appropriate.",
            "Repository uses t-test (R/09) and distribution/robust risk analyses instead.",
        ],
        font_size=17,
    )

    # Sources slide
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Sources used in this deck"
    tf = slide.shapes.placeholders[1].text_frame
    tf.clear()
    sources = [
        "Scripts: R/03_data_cleaning.R, R/07_distribution_fitting.R, R/09_hypothesis_testing.R, R/14_var_analysis.R",
        "Tables: output/tables/cleaning_summary.csv, 5stock_descriptive_stats.csv, model_comparison_summary.csv,",
        "        goodness_of_fit_summary.csv, monte_carlo_summary.csv, clt_summary.csv, correlation_matrix.csv, t_fit_params.csv",
        "Figures: output/figures/historical_prices.png, extreme_return_plot.png, qq_plots_combined.png,",
        "         correlation_heatmap.png, monte_carlo_simulation.png",
    ]
    for i, s in enumerate(sources):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = s
        p.font.size = Pt(16)
        p.alignment = PP_ALIGN.LEFT

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    prs.save(str(PPT_PATH))
    print(f"Generated: {PPT_PATH}")


if __name__ == "__main__":
    main()
